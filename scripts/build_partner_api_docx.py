from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path("docs/PARTNER-API-v1.md")
OUTPUT = Path(".docs/PARTNER-API-v1.docx")
PAGE_WIDTH_DXA = 9360
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
TABLE_FILL = "E8EEF5"


def set_run_font(run, size=11, bold=None, color="000000", italic=None, name="Calibri"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    width = tc_pr.find(qn("w:tcW"))
    if width is None:
        width = OxmlElement("w:tcW")
        tc_pr.append(width)
    width.set(qn("w:w"), str(width_dxa))
    width.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_layout = tbl_pr.first_child_found_in("w:tblLayout")
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths):
        col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            margins = tc_pr.first_child_found_in("w:tcMar")
            if margins is None:
                margins = OxmlElement("w:tcMar")
                tc_pr.append(margins)
            for side, value in (("top", 80), ("bottom", 80), ("start", 120), ("end", 120)):
                node = margins.find(qn(f"w:{side}"))
                if node is None:
                    node = OxmlElement(f"w:{side}")
                    margins.append(node)
                node.set(qn("w:w"), str(value))
                node.set(qn("w:type"), "dxa")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, size=9, color="666666")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def add_inline_markup(paragraph, text, base_size=11, code=False):
    pattern = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*)")
    for token in pattern.split(text):
        if not token:
            continue
        if token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=base_size - 1, color="1F4D78", name="Consolas")
        elif token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=base_size, bold=True)
        elif token.startswith("*") and token.endswith("*"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=base_size, italic=True)
        else:
            run = paragraph.add_run(token)
            set_run_font(run, size=base_size, name="Consolas" if code else "Calibri")


def add_table(doc, rows):
    cells = [[cell.strip() for cell in row.strip().strip("|").split("|")] for row in rows]
    cells = [row for row in cells if not all(re.fullmatch(r"\s*:?-{3,}:?\s*", value) for value in row)]
    if not cells:
        return
    column_count = len(cells[0])
    table = doc.add_table(rows=0, cols=column_count)
    max_lengths = [max(len(row[index]) if index < len(row) else 0 for row in cells) for index in range(column_count)]
    weighted = [max(12, length) for length in max_lengths]
    total = sum(weighted)
    widths = [max(1000, int(PAGE_WIDTH_DXA * item / total)) for item in weighted]
    widths[-1] += PAGE_WIDTH_DXA - sum(widths)
    for row_index, values in enumerate(cells):
        row = table.add_row()
        for index, value in enumerate(values):
            cell = row.cells[index]
            if row_index == 0:
                set_cell_shading(cell, TABLE_FILL)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            add_inline_markup(paragraph, value, base_size=9)
            for run in paragraph.runs:
                run.bold = row_index == 0
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for style_name, size, color, before, after in (("Heading 1", 16, BLUE, 18, 10), ("Heading 2", 13, BLUE, 14, 7), ("Heading 3", 12, DARK_BLUE, 10, 5)):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    footer = section.footer.paragraphs[0]
    add_page_number(footer)


def main():
    if not SOURCE.exists():
        raise SystemExit(f"Missing source: {SOURCE}")
    OUTPUT.parent.mkdir(exist_ok=True)
    doc = Document()
    configure_document(doc)
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        if line.startswith("```"):
            language = line[3:].strip()
            index += 1
            block = []
            while index < len(lines) and not lines[index].startswith("```"):
                block.append(lines[index])
                index += 1
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(6)
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "F4F6F9")
            paragraph._p.get_or_add_pPr().append(shade)
            run = paragraph.add_run("\n".join(block))
            set_run_font(run, size=8.5, color="1F2937", name="Consolas")
            index += 1
            continue
        if line.startswith("|") and "|" in line[1:]:
            table_lines = []
            while index < len(lines) and lines[index].startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_table(doc, table_lines)
            continue
        match = re.match(r"^(#{1,3})\s+(.*)$", line)
        if match:
            level = len(match.group(1))
            text = match.group(2)
            if level == 1:
                paragraph = doc.add_paragraph()
                paragraph.paragraph_format.space_after = Pt(8)
                run = paragraph.add_run(text)
                set_run_font(run, size=23, bold=True, color="0B2545")
            else:
                paragraph = doc.add_paragraph(style=f"Heading {level - 1}")
                add_inline_markup(paragraph, text, base_size={2: 16, 3: 13}[level])
            index += 1
            continue
        if not line.strip():
            index += 1
            continue
        bullet = re.match(r"^[-*]\s+(.*)$", line)
        numbered = re.match(r"^\d+\.\s+(.*)$", line)
        if bullet or numbered:
            style = "List Bullet" if bullet else "List Number"
            paragraph = doc.add_paragraph(style=style)
            paragraph.paragraph_format.space_after = Pt(4)
            add_inline_markup(paragraph, (bullet or numbered).group(1))
            index += 1
            continue
        paragraph = doc.add_paragraph()
        add_inline_markup(paragraph, line.rstrip())
        index += 1
    doc.core_properties.title = "Visa Compass Partner API v1"
    doc.core_properties.subject = "Partner API integration specification"
    doc.core_properties.author = "Visa Compass"
    doc.save(OUTPUT)
    print(OUTPUT.resolve())


if __name__ == "__main__":
    main()
